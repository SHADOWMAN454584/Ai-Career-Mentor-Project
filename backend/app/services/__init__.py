import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..catalog import ROLES
from ..llm import extract_resume_with_gemini, reason_about_profile
from ..models import AnalysisJob, AnalysisResult, CandidateProfile, Resume
from .skill_analysis import (
    KNOWN_SKILLS,
    build_roadmap,
    detect_gaps,
    normalize_skill,
    readiness,
    recommend_courses,
    role_scores,
    similarity,
)

logger = logging.getLogger(__name__)

GENERIC_SKILL_TOKENS = {"projects", "internship"}
PROJECT_HEADING_RE = re.compile(
    r"^\s*(?:projects?|personal projects?|academic projects?|selected projects?)\s*[:\-]?\s*$",
    re.IGNORECASE,
)
SECTION_HEADING_RE = re.compile(
    r"^\s*(?:education|experience|work experience|professional experience|internships?|"
    r"skills|technical skills|certifications?|achievements?|awards?|leadership|"
    r"activities|summary|profile|contact)\b",
    re.IGNORECASE,
)

__all__ = [
    "answer_similarity",
    "build_roadmap",
    "detect_gaps",
    "normalize_skill",
    "parse_resume",
    "readiness",
    "recommend_courses",
    "role_scores",
    "run_analysis",
    "similarity",
]


def extract_text(path: str) -> str:
    """Robust multi-engine document text extraction for PDF and DOCX files."""
    source = Path(path)
    if not source.exists():
        raise HTTPException(status_code=404, detail=f"Resume file not found on disk: {source.name}")

    suffix = source.suffix.lower()

    if suffix == ".pdf":
        extracted_text = ""

        # Primary engine: PyMuPDF (fitz)
        try:
            import fitz
            with fitz.open(source) as pdf:
                pages_text = []
                for page in pdf:
                    ptxt = page.get_text()
                    if ptxt:
                        pages_text.append(ptxt)
                extracted_text = "\n".join(pages_text).strip()
        except Exception as fitz_err:
            logger.warning("PyMuPDF failed to parse %s: %s", source.name, fitz_err)

        # Fallback engine 1: pypdf
        if not extracted_text:
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(source))
                pages_text = []
                for page in reader.pages:
                    ptxt = page.extract_text()
                    if ptxt:
                        pages_text.append(ptxt)
                extracted_text = "\n".join(pages_text).strip()
            except Exception as pypdf_err:
                logger.warning("pypdf fallback failed to parse %s: %s", source.name, pypdf_err)

        # Fallback engine 2: pdfplumber
        if not extracted_text:
            try:
                import pdfplumber
                with pdfplumber.open(source) as pdf:
                    pages_text = []
                    for page in pdf.pages:
                        ptxt = page.extract_text()
                        if ptxt:
                            pages_text.append(ptxt)
                    extracted_text = "\n".join(pages_text).strip()
            except Exception as plumb_err:
                logger.warning("pdfplumber fallback failed to parse %s: %s", source.name, plumb_err)

        # Validate that actual character text was ingested
        cleaned = re.sub(r"\s+", " ", extracted_text).strip()
        if not cleaned or len(cleaned) < 15:
            raise HTTPException(
                status_code=422,
                detail="Unable to extract text from the PDF. The file may be empty, image-only (scanned), or encrypted. Please upload a standard text-based PDF or DOCX resume.",
            )
        return extracted_text

    if suffix == ".docx":
        try:
            from docx import Document
            doc = Document(source)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            extracted_text = "\n".join(paragraphs).strip()
            if not extracted_text:
                raise ValueError("DOCX document has no readable text")
            return extracted_text
        except Exception as docx_err:
            raise HTTPException(status_code=422, detail=f"Could not read DOCX document: {docx_err}") from docx_err

    raise HTTPException(status_code=415, detail="Only PDF and DOCX resumes are supported")


def _extract_projects(text: str) -> list[str]:
    """Extract project entries without truncating at decimal points or periods."""
    lines = [line.strip(" \t•-–—") for line in text.splitlines()]
    items: list[str] = []
    in_projects = False

    for line in lines:
        if not line:
            continue
        if PROJECT_HEADING_RE.match(line):
            in_projects = True
            continue
        if in_projects and SECTION_HEADING_RE.match(line):
            break
        if in_projects:
            items.append(line)

    if not items:
        for line in lines:
            if len(line) >= 12 and re.search(r"\b(?:built|developed|created|designed|implemented|launched)\b", line, re.IGNORECASE):
                items.append(line)

    cleaned: list[str] = []
    for item in items:
        item = re.sub(r"^[\-•*]\s*", "", item).strip()
        if len(item) < 4:
            continue
        if item.lower() in {"projects", "personal projects", "academic projects"}:
            continue
        cleaned.append(item[:180])
        if len(cleaned) >= 8:
            break
    return cleaned


def _parse_resume_regex(raw_text: str) -> dict[str, Any]:
    """Deterministic catalog-matching and regex extraction for resumes."""
    text = re.sub(r"[\t ]+", " ", raw_text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lower = text.lower()

    skills = sorted(
        {
            normalize_skill(skill)
            for skill in KNOWN_SKILLS
            if re.search(rf"(?<!\w){re.escape(skill)}(?!\w)", lower)
            and normalize_skill(skill) not in GENERIC_SKILL_TOKENS
        }
    )
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone_source = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", " ", text)
    phone_source = re.sub(r"https?://\S+|www\.\S+", " ", phone_source)
    phone = re.search(
        r"(?<![\w])(?:\+\d{1,3}[\s.-]?)?(?:\(\d{2,4}\)[\s.-]?)?\d{3}[\s.-]?\d{3,4}(?:[\s.-]?\d{4})?(?![\w])",
        phone_source,
    )
    cgpa = re.search(r"(?:cgpa|gpa)\s*[:=-]?\s*(\d(?:\.\d{1,2})?)", lower)
    year = re.search(r"(?:graduat(?:ion|ing)|batch|class\s*of)?\s*[:=-]?\s*(20[12]\d)\b", lower)
    experience_years = re.search(r"(\d+(?:\.\d+)?)\s*(?:years?|yrs?)\s+(?:of\s+)?experience", lower)
    projects = _extract_projects(text)
    degree_match = re.search(r"\b(B\.?Tech[^\n,]*|M\.?Tech[^\n,]*|BCA|MCA|B\.?Sc|M\.?Sc|Bachelor[^\n,]*|Master[^\n,]*)", text, flags=re.I)

    return {
        "skills": skills,
        "education": {
            "degree": degree_match.group(0).strip() if degree_match else None,
            "graduation_year": year.group(1) if year else None,
            "cgpa": float(cgpa.group(1)) if cgpa else None,
        },
        "experience": {"years": float(experience_years.group(1)) if experience_years else 0.0},
        "projects": [p.strip() for p in projects[:8]],
        "contacts": {"email": email.group(0) if email else None, "phone": phone.group(0) if phone else None},
    }


def parse_resume(raw_text: str) -> dict[str, Any]:
    """High-fidelity resume parsing combining Gemini intelligence with offline deterministic matching."""
    # 1. Deterministic baseline
    local_parsed = _parse_resume_regex(raw_text)

    # 2. Intelligent Gemini parsing
    gemini_parsed = extract_resume_with_gemini(raw_text)

    if gemini_parsed:
        # Merge skills: Gemini skills normalized + detected known catalog skills
        combined_skills = set(local_parsed.get("skills", []))
        for skill in gemini_parsed.get("skills", []):
            norm = normalize_skill(skill)
            if norm and norm not in GENERIC_SKILL_TOKENS:
                combined_skills.add(norm)

        # Merge education
        education = gemini_parsed.get("education") or {}
        if not education.get("cgpa") and local_parsed["education"].get("cgpa"):
            education["cgpa"] = local_parsed["education"]["cgpa"]
        if not education.get("graduation_year") and local_parsed["education"].get("graduation_year"):
            education["graduation_year"] = local_parsed["education"]["graduation_year"]
        if not education.get("degree") and local_parsed["education"].get("degree"):
            education["degree"] = local_parsed["education"]["degree"]

        # Merge projects and contacts
        projects = gemini_parsed.get("projects") or local_parsed.get("projects", [])
        contacts = gemini_parsed.get("contacts") or local_parsed.get("contacts", {})
        experience = gemini_parsed.get("experience") or local_parsed.get("experience", {"years": 0.0})

        return {
            "skills": sorted(combined_skills),
            "education": education,
            "experience": experience,
            "projects": projects[:8],
            "contacts": contacts,
        }

    return local_parsed


def save_profile(db: Session, resume: Resume, parsed: dict[str, Any]) -> CandidateProfile:
    profile = db.scalar(select(CandidateProfile).where(CandidateProfile.resume_id == resume.id))
    if not profile:
        profile = CandidateProfile(user_id=resume.user_id, resume_id=resume.id)
        db.add(profile)
    profile.skills = parsed["skills"]
    profile.education = parsed["education"]
    profile.experience = parsed["experience"]
    profile.projects = parsed["projects"]
    profile.contacts = parsed["contacts"]
    profile.extracted_at = datetime.utcnow()
    return profile


def run_analysis(db: Session, job: AnalysisJob) -> AnalysisResult:
    job.status, job.started_at = "running", datetime.utcnow()
    db.commit()
    resume = db.get(Resume, job.resume_id)
    try:
        resume.raw_text = extract_text(resume.file_path)
        parsed = parse_resume(resume.raw_text)
        save_profile(db, resume, parsed)
        reasoning = reason_about_profile(parsed, job.target_role, ROLES[job.target_role], job.timeline_weeks)
        if reasoning:
            gaps = reasoning.skill_gaps
            probability, factors, roadmap = reasoning.placement_probability, reasoning.shap_factors, reasoning.roadmap
            model_version = f"llm-{reasoning.provider}"
            for week in roadmap:
                for task in week["tasks"]:
                    task["task_id"] = str(uuid4())
        else:
            gaps = detect_gaps(parsed["skills"], job.target_role)
            probability, factors = readiness(parsed)
            roadmap = build_roadmap(gaps, job.timeline_weeks)
            model_version = "demo-readiness-v1"
        result = AnalysisResult(job_id=job.id, user_id=job.user_id, role_scores=role_scores(parsed["skills"]), skill_gaps=gaps, courses=recommend_courses(gaps), placement_probability=probability, shap_factors=factors, roadmap=roadmap, model_version=model_version)
        db.add(result)
        job.status, job.completed_at = "completed", datetime.utcnow()
        db.commit()
        db.refresh(result)
        return result
    except Exception as exc:
        job.status, job.error, job.completed_at = "failed", str(exc), datetime.utcnow()
        db.commit()
        raise


def answer_similarity(answer: str, model_answer: str) -> float:
    answer_words = set(re.findall(r"\w+", answer.lower()))
    model_words = set(re.findall(r"\w+", model_answer.lower()))
    return round(100 * len(answer_words & model_words) / max(1, len(answer_words | model_words)), 1)
